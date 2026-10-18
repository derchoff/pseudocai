"""Statistiques de traitement, lues dans le container Cosmos DB « statistiques » : affichage à
l'écran et classeur Excel.

    uv run --python 3.12 --with azure-cosmos --with azure-identity --with openpyxl \\
        python scripts/statistiques.py --cosmos cosmos-pseudo-dev-ab12 --debut 2026-10-01 --fin 2026-10-31

Une fiche par traitement (document, pièce de mail, JSON de métadonnées ; voir src/pseudo/statistics/statistics.py).
Les périodes suivent la date de DÉPOSE dans incoming/ (heure de Paris). Calcule, sur la période et
par jour, mois ou année (--par) :
- documents déposés (un mail compte pour un), pages traitées et pages par document, pièces de mail
  et JSON compris, documents publiés ou en quarantaine ;
- pourcentage de mise en quarantaine ;
- répartition par type : PDF hors mail, image hors mail (par type MIME), mail, PDF dans un mail,
  image dans un mail, autre pièce d'un mail, fichier refusé au routage, JSON de métadonnées ;
- tours de vérification (1, 2, 3… tours ; résidus après le dernier tour) ;
- pages réellement passées par l'OCR et par les agents IA ;
- consommation (appels et jetons des agents IA, enregistrements Azure AI Language)
  et, avec les prix (--prix-*), coût estimé et coût par page.

Un document retraité ou redéposé compte à chaque traitement : il consomme à nouveau OCR et IA.
Connexion : identité Azure de la session (az login), rôle « Cosmos DB Built-in Data Reader » sur le
container (variable Terraform statistics_reader_principal_ids). --fiches lit un export JSON à la place.
Les fiches ne contiennent aucune donnée personnelle (ni nom de fichier, ni valeur).
"""

from __future__ import annotations

import argparse
import json
import statistics as stats
import sys
from collections import defaultdict
from datetime import date
from pathlib import Path

CATEGORIES = {
    "pdf": "PDF hors mail",
    "image": "Image hors mail",
    "mail": "Mail (.eml)",
    "pdf_mail": "PDF dans un mail",
    "image_mail": "Image dans un mail",
    "autre_mail": "Autre pièce dans un mail",
    "autre": "Fichier refusé au routage",
    "json": "JSON de métadonnées GED",
}
CATEGORY_ORDER = list(CATEGORIES)
AGENTS = ("detection", "verification", "classification")
AGENT_LABELS = {"detection": "Détection", "verification": "Vérification", "classification": "Classification"}


# ------------------------------------------------------------------------- lecture des fiches


def cosmos_endpoint(value: str) -> str:
    """Nom du compte (cosmos-pseudo-dev-ab12) ou URL complète -> URL du compte."""
    if value.startswith("https://"):
        return value
    return f"https://{value}.documents.azure.com:443/"


def load_from_cosmos(endpoint: str, database: str, container: str, start: str, end: str) -> list[dict]:
    """Fiches de la période [start, end] (jours de dépose, AAAA-MM-JJ), lues avec l'identité Azure."""
    from azure.cosmos import CosmosClient
    from azure.identity import DefaultAzureCredential

    client = CosmosClient(endpoint, credential=DefaultAzureCredential())
    items = client.get_database_client(database).get_container_client(container).query_items(
        query="SELECT * FROM c WHERE c.jour >= @debut AND c.jour <= @fin",
        parameters=[{"name": "@debut", "value": start}, {"name": "@fin", "value": end}],
        enable_cross_partition_query=True,
    )
    return [{k: v for k, v in item.items() if not k.startswith("_")} for item in items]


def load_from_file(path: Path, start: str, end: str) -> list[dict]:
    """Fiches d'un export JSON (liste de fiches), filtrées sur la période."""
    rows = json.loads(path.read_text(encoding="utf-8-sig"))
    return [r for r in rows if start <= r.get("jour", "") <= end]


# ------------------------------------------------------------------------- calculs


def usage(row: dict) -> dict:
    """Consommation à plat d'une fiche : pages OCR, Language, appels et jetons (total et par agent)."""
    c = row.get("consommation") or {}
    ia = c.get("ia") or {}
    out = {
        "pages_ocr": int(c.get("pages_ocr", 0) or 0),
        "pages_ia": int(c.get("pages_ia", 0) or 0),
        "language": int((c.get("language") or {}).get("enregistrements", 0) or 0),
        "appels": 0, "entree": 0, "sortie": 0, "raisonnement": 0,
    }
    for agent in AGENTS:
        a = ia.get(agent) or {}
        values = (int(a.get("appels", 0) or 0), int(a.get("jetons_entree", 0) or 0),
                  int(a.get("jetons_sortie", 0) or 0), int(a.get("jetons_raisonnement", 0) or 0))
        out[agent] = values
        out["appels"] += values[0]
        out["entree"] += values[1]
        out["sortie"] += values[2]
        out["raisonnement"] += values[3]
    return out


def cost(u: dict, prices: dict) -> float | None:
    """Coût estimé d'une consommation, ou None sans prix. Les jetons de sortie comprennent déjà le
    raisonnement ; les jetons d'entrée sont comptés au plein tarif (le cache n'est pas déduit)."""
    if not any(prices.values()):
        return None
    return (
        u["pages_ocr"] / 1000 * prices["ocr"]
        + u["entree"] / 1_000_000 * prices["entree"]
        + u["sortie"] / 1_000_000 * prices["sortie"]
        + u["language"] / 1000 * prices["language"]
    )


def period_of(row: dict, by: str) -> str:
    """Clé de période d'une fiche : AAAA-MM-JJ, AAAA-MM ou AAAA."""
    return {"jour": row["jour"], "mois": row["jour"][:7], "annee": row["jour"][:4]}[by]


def _pct(part: int, whole: int) -> float | None:
    """Pourcentage arrondi, ou None si la base est nulle."""
    return round(100 * part / whole, 1) if whole else None


def _add(total: dict, u: dict) -> None:
    """Ajoute une consommation à plat à un total."""
    for k in ("pages_ocr", "pages_ia", "language", "appels", "entree", "sortie", "raisonnement"):
        total[k] = total.get(k, 0) + u[k]


def documents(rows: list[dict], prices: dict) -> list[dict]:
    """Une ligne par document déposé : ses pages et sa consommation, pièces de mail et JSON compris."""
    children: dict[str, list[dict]] = defaultdict(list)
    for r in rows:
        if r.get("niveau") == "piece" and r.get("mailJobId"):
            children[r["mailJobId"]].append(r)
    json_by_doc: dict[str, list[dict]] = defaultdict(list)
    for r in rows:
        if r.get("niveau") == "annexe":
            json_by_doc[r["documentId"]].append(r)
    out = []
    for r in rows:
        if r.get("niveau") != "document":
            continue
        parts = children.get(r["jobId"], [])
        total: dict = {}
        for item in [r, *parts, *json_by_doc.get(r["documentId"], [])]:
            _add(total, usage(item))
        out.append({
            "jour": r["jour"], "heure": r.get("hour"), "documentId": r["documentId"],
            "categorie": r["categorie"], "mime": r.get("mime"), "statut": r["statut"],
            "motifs": ", ".join(r.get("motifs") or []), "domain": r.get("domain"),
            "pages": int(r.get("pages", 0) or 0) + sum(int(p.get("pages", 0) or 0) for p in parts),
            "pieces": len(parts), "tours": r.get("toursVerification", 0), **total, "cout": cost(total, prices),
        })
    return sorted(out, key=lambda d: (d["jour"], d["heure"] or 0, d["documentId"]))


def compute(rows: list[dict], by: str, prices: dict) -> dict:
    """Tous les tableaux : synthèse, par période, par type, tours de vérification, agents, documents."""
    docs = documents(rows, prices)
    total: dict = {}
    for r in rows:
        _add(total, usage(r))
    pages = sum(int(r.get("pages", 0) or 0) for r in rows)
    finished = [d for d in docs if d["statut"] in ("ok", "quarantine")]
    quarantined = [d for d in finished if d["statut"] == "quarantine"]
    parts = [r for r in rows if r.get("niveau") == "piece"]
    reprocessed = sum(1 for n in _count_by(docs, "documentId").values() if n > 1)
    total_cost = cost(total, prices)
    summary = [
        ("Fiches (traitements)", len(rows)),
        ("Documents déposés (un mail = 1)", len(docs)),
        ("dont mails", sum(1 for d in docs if d["categorie"] == "mail")),
        ("Documents publiés", len(finished) - len(quarantined)),
        ("Documents en quarantaine", len(quarantined)),
        ("% de mise en quarantaine", _pct(len(quarantined), len(finished))),
        ("Documents traités plus d'une fois", reprocessed),
        ("Pièces de mails traitées", sum(1 for p in parts if p["statut"] in ("ok", "quarantine"))),
        ("Pièces de mails en quarantaine", sum(1 for p in parts if p["statut"] == "quarantine")),
        ("Pièces refusées (non prises en charge)", sum(1 for p in parts if p["statut"] == "refusee")),
        ("Images retirées des mails (GIF, logos)", sum(1 for p in parts if p["statut"] == "retiree")),
        ("JSON de métadonnées", sum(1 for r in rows if r.get("niveau") == "annexe")),
        ("Pages traitées (documents et pièces)", pages),
        ("Pages par document (moyenne)", round(stats.mean(d["pages"] for d in docs), 2) if docs else None),
        ("Pages par document (médiane)", stats.median(d["pages"] for d in docs) if docs else None),
        ("Pages par document (maximum)", max((d["pages"] for d in docs), default=None)),
        ("Pages passées par l'OCR (Document Intelligence)", total.get("pages_ocr", 0)),
        ("Pages passées par les agents IA", total.get("pages_ia", 0)),
        ("Textes passés par les agents IA (corps de mail, JSON)",
         sum(1 for r in rows if r.get("categorie") in ("mail", "json") and usage(r)["appels"])),
        ("Appels aux agents IA", total.get("appels", 0)),
        ("Jetons d'entrée", total.get("entree", 0)),
        ("Jetons de sortie (raisonnement compris)", total.get("sortie", 0)),
        ("dont jetons de raisonnement", total.get("raisonnement", 0)),
        ("Enregistrements Azure AI Language", total.get("language", 0)),
        ("Coût estimé", _money(total_cost)),
        ("Coût par page traitée", _money(total_cost / pages if total_cost is not None and pages else None, 4)),
    ]
    return {
        "Synthèse": (["Indicateur", "Valeur"], [list(s) for s in summary]),
        "Par période": _by_period(rows, docs, by, prices),
        "Par type": _by_type(rows, prices),
        "Tours de vérification": _rounds(rows, by),
        "Agents IA": _agents(rows, prices),
        "Documents": (
            ["Jour", "Heure", "Document (empreinte)", "Catégorie", "Type MIME", "Statut", "Motifs", "Domaine",
             "Pages", "Pièces", "Tours", "Pages OCR", "Pages IA", "Appels IA", "Jetons entrée", "Jetons sortie", "Coût"],
            [[d["jour"], d["heure"], d["documentId"], CATEGORIES.get(d["categorie"], d["categorie"]), d["mime"],
              d["statut"], d["motifs"], d["domain"], d["pages"], d["pieces"], d["tours"], d.get("pages_ocr", 0),
              d.get("pages_ia", 0), d.get("appels", 0), d.get("entree", 0), d.get("sortie", 0), _money(d["cout"], 4)] for d in docs],
        ),
    }


def _count_by(items: list[dict], key: str) -> dict:
    """Nombre d'éléments par valeur de la clé."""
    out: dict = defaultdict(int)
    for item in items:
        out[item[key]] += 1
    return out


def _by_period(rows: list[dict], docs: list[dict], by: str, prices: dict) -> tuple[list, list]:
    """Documents, pages, quarantaine et consommation par période."""
    groups: dict[str, dict] = defaultdict(lambda: {"docs": 0, "mails": 0, "fin": 0, "quar": 0, "pages": 0, "u": {}})
    for d in docs:
        g = groups[period_of(d, by)]
        g["docs"] += 1
        g["mails"] += d["categorie"] == "mail"
        g["fin"] += d["statut"] in ("ok", "quarantine")
        g["quar"] += d["statut"] == "quarantine"
    for r in rows:
        g = groups[period_of(r, by)]
        g["pages"] += int(r.get("pages", 0) or 0)
        _add(g["u"], usage(r))
    out = []
    for key in sorted(groups):
        g = groups[key]
        c = cost(g["u"], prices) if g["u"] else None
        out.append([key, g["docs"], g["mails"], g["quar"], _pct(g["quar"], g["fin"]), g["pages"],
                    g["u"].get("pages_ocr", 0), g["u"].get("pages_ia", 0), g["u"].get("entree", 0), g["u"].get("sortie", 0),
                    _money(c), _money(c / g["pages"] if c is not None and g["pages"] else None, 4)])
    headers = ["Période", "Documents", "dont mails", "En quarantaine", "% quarantaine", "Pages", "Pages OCR",
               "Pages IA", "Jetons entrée", "Jetons sortie", "Coût", "Coût / page"]
    return headers, out


def _by_type(rows: list[dict], prices: dict) -> tuple[list, list]:
    """Par catégorie et type MIME : nombre, statuts, pages, consommation."""
    groups: dict[tuple, dict] = defaultdict(lambda: defaultdict(int))
    usages: dict[tuple, dict] = defaultdict(dict)
    for r in rows:
        key = (r["categorie"], r.get("mime") or "")
        g = groups[key]
        g["n"] += 1
        g[r["statut"]] += 1
        g["pages"] += int(r.get("pages", 0) or 0)
        _add(usages[key], usage(r))
    out = []
    for key in sorted(groups, key=lambda k: (CATEGORY_ORDER.index(k[0]) if k[0] in CATEGORY_ORDER else 99, k[1])):
        g, u = groups[key], usages[key]
        out.append([CATEGORIES.get(key[0], key[0]), key[1], g["n"], g["ok"], g["quarantine"], g["refusee"],
                    g["retiree"], g["non_traitee"], _pct(g["quarantine"], g["ok"] + g["quarantine"]), g["pages"],
                    u.get("pages_ocr", 0), u.get("pages_ia", 0), u.get("entree", 0) + u.get("sortie", 0),
                    _money(cost(u, prices))])
    headers = ["Catégorie", "Type MIME", "Nombre", "Publiés", "Quarantaine", "Refusées", "Retirées", "Non traitées",
               "% quarantaine", "Pages", "Pages OCR", "Pages IA", "Jetons", "Coût"]
    return headers, out


def _rounds(rows: list[dict], by: str) -> tuple[list, list]:
    """Tours de vérification par période, sur les textes vérifiés (documents, mails, pièces, JSON)."""
    checked = [r for r in rows if int(r.get("toursVerification", 0) or 0) > 0]
    groups: dict[str, list[dict]] = defaultdict(list)
    for r in checked:
        groups[period_of(r, by)].append(r)
    out = []
    for key in sorted(groups):
        g = groups[key]
        tours = [int(r["toursVerification"]) for r in g]
        residues = sum(1 for r in g if r.get("residus"))
        out.append([key, len(g), tours.count(1), tours.count(2), tours.count(3), sum(1 for t in tours if t >= 4),
                    round(stats.mean(tours), 2), max(tours), residues, _pct(residues, len(g))])
    headers = ["Période", "Textes vérifiés", "1 tour", "2 tours", "3 tours", "4 tours et plus", "Tours (moyenne)",
               "Tours (maximum)", "Résidus après le dernier tour", "% résidus"]
    return headers, out


def _agents(rows: list[dict], prices: dict) -> tuple[list, list]:
    """Appels et jetons par agent IA."""
    out = []
    for agent in AGENTS:
        calls = sum(usage(r)[agent][0] for r in rows)
        tokens_in = sum(usage(r)[agent][1] for r in rows)
        tokens_out = sum(usage(r)[agent][2] for r in rows)
        reasoning = sum(usage(r)[agent][3] for r in rows)
        u = {"pages_ocr": 0, "language": 0, "entree": tokens_in, "sortie": tokens_out}
        out.append([AGENT_LABELS[agent], calls, tokens_in, tokens_out, reasoning, _money(cost(u, prices))])
    return ["Agent", "Appels", "Jetons entrée", "Jetons sortie", "dont raisonnement", "Coût"], out


def _money(value: float | None, digits: int = 2) -> float | None:
    """Montant arrondi, ou None."""
    return None if value is None else round(value, digits)


# ------------------------------------------------------------------------- sorties


def _fmt(value) -> str:
    """Valeur affichée à l'écran : séparateur de milliers, vide pour None."""
    if value is None:
        return "—"
    if isinstance(value, bool):
        return "oui" if value else "non"
    if isinstance(value, int):
        return f"{value:,}".replace(",", " ")
    if isinstance(value, float):
        return f"{value:,.4f}".rstrip("0").rstrip(".").replace(",", " ")
    return str(value)


def print_table(title: str, headers: list[str], rows: list[list], limit: int | None = None) -> None:
    """Tableau aligné dans le terminal."""
    print(f"\n{title}\n{'=' * len(title)}")
    if not rows:
        print("(aucune donnée)")
        return
    shown = rows[:limit] if limit else rows
    cells = [[_fmt(v) for v in r] for r in shown]
    widths = [max(len(h), *(len(c[i]) for c in cells)) for i, h in enumerate(headers)]
    print("  ".join(h.ljust(w) for h, w in zip(headers, widths)))
    print("  ".join("-" * w for w in widths))
    for c in cells:
        print("  ".join(v.rjust(w) if v[:1].isdigit() or v[:1] == "—" else v.ljust(w) for v, w in zip(c, widths)))
    if limit and len(rows) > limit:
        print(f"… {len(rows) - limit} ligne(s) de plus dans le fichier Excel")


def write_excel(path: Path, tables: dict, context: list[tuple[str, str]]) -> None:
    """Classeur : un onglet par tableau, en-têtes en couleur, première ligne figée."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    font = Font(name="Arial", size=10)
    wb = Workbook()
    wb.remove(wb.active)
    for name, (headers, rows) in tables.items():
        ws = wb.create_sheet(name[:31])
        start = 1
        if name == "Synthèse":
            for label, value in context:
                ws.append([label, value])
            ws.append([])
            start = len(context) + 2
        ws.append(headers)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(start, c)
            cell.font, cell.fill = header_font, header_fill
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=1):
            for cell in row:
                if cell.row != start:
                    cell.font = font
        for c, h in enumerate(headers, start=1):
            values = [h] + [r[c - 1] for r in rows]
            ws.column_dimensions[get_column_letter(c)].width = min(60, max(10, *(len(str(v or "")) + 2 for v in values)))
        ws.freeze_panes = ws.cell(start + 1, 1)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    """Arguments de la ligne de commande."""
    today = date.today()
    p = argparse.ArgumentParser(description="Statistiques de traitement (container Cosmos DB « statistiques »).")
    source = p.add_mutually_exclusive_group(required=True)
    source.add_argument("--cosmos", help="compte Cosmos DB : nom (cosmos-pseudo-dev-ab12) ou URL https://…")
    source.add_argument("--fiches", type=Path, help="export JSON des fiches (liste), à la place de Cosmos DB")
    p.add_argument("--base", default="pseudonymisation", help="base Cosmos DB (défaut : pseudonymisation)")
    p.add_argument("--container", default="statistiques", help="container des fiches (défaut : statistiques)")
    p.add_argument("--debut", default=today.replace(day=1).isoformat(), help="premier jour de dépose, AAAA-MM-JJ (défaut : 1er du mois)")
    p.add_argument("--fin", default=today.isoformat(), help="dernier jour de dépose, AAAA-MM-JJ (défaut : aujourd'hui)")
    p.add_argument("--par", choices=("jour", "mois", "annee"), default="jour", help="regroupement des tableaux par période")
    p.add_argument("--excel", type=Path, help="classeur produit (défaut : statistiques_<debut>_<fin>.xlsx)")
    p.add_argument("--sans-excel", action="store_true", help="affichage à l'écran seulement")
    p.add_argument("--prix-ocr", type=float, default=0.0, help="prix de 1 000 pages Document Intelligence")
    p.add_argument("--prix-entree", type=float, default=0.0, help="prix de 1 million de jetons d'entrée")
    p.add_argument("--prix-sortie", type=float, default=0.0, help="prix de 1 million de jetons de sortie")
    p.add_argument("--prix-language", type=float, default=0.0, help="prix de 1 000 enregistrements Azure AI Language")
    args = p.parse_args(argv)
    for value in (args.debut, args.fin):
        date.fromisoformat(value)  # lève ValueError si la date est mal écrite
    return args


def main(argv: list[str] | None = None) -> int:
    """Lit les fiches de la période, affiche les tableaux et écrit le classeur Excel."""
    args = parse_args(argv)
    prices = {"ocr": args.prix_ocr, "entree": args.prix_entree, "sortie": args.prix_sortie, "language": args.prix_language}
    if args.fiches:
        rows = load_from_file(args.fiches, args.debut, args.fin)
        origin = str(args.fiches)
    else:
        origin = cosmos_endpoint(args.cosmos)
        rows = load_from_cosmos(origin, args.base, args.container, args.debut, args.fin)
    tables = compute(rows, args.par, prices)
    context = [("Source", origin), ("Période de dépose", f"{args.debut} → {args.fin}"), ("Regroupement", args.par),
               ("Prix", ", ".join(f"{k} = {v}" for k, v in prices.items()) if any(prices.values()) else "non renseignés")]
    print(f"Statistiques — dépôts du {args.debut} au {args.fin} ({len(rows)} fiche(s), source : {origin})")
    for name, (headers, table_rows) in tables.items():
        if name == "Documents":
            continue  # une ligne par document : Excel seulement
        print_table(name, headers, table_rows, limit=40)
    if not args.sans_excel:
        path = args.excel or Path(f"statistiques_{args.debut}_{args.fin}.xlsx")
        write_excel(path, tables, context)
        print(f"\nClasseur Excel : {path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
