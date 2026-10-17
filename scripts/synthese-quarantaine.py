"""Synthèse Excel des documents en quarantaine, à partir d'une copie locale du container quarantine.

    uv run --python 3.12 --with openpyxl python scripts/synthese-quarantaine.py echantillon/quarantaine

Parcourt tous les sous-dossiers, lit chaque motif.json (mail) et *.motif.json (document), puis écrit
<dossier>/synthese-quarantaine.xlsx :
- Synthèse : nombre de mails et de documents par motif, et par dossier de type ;
- Détail : une ligne par fichier en quarantaine ;
- Pièces non prises en charge : pour chaque pièce refusée d'un mail, sa structure MIME
  (type, disposition, nom, taille, format réel reconnu par ses premiers octets) ;
- OCR faible : confiance OCR de chaque page lue par OCR, la plus basse d'abord. Les pièces d'un mail
  concernées sont extraites dans <dossier du mail>/pieces_ocr_faible/ pour être ouvertes directement ;
- Échecs de localisation : chaque valeur signalée par un agent mais introuvable dans le texte
  (diagnostic.localisation des rapports), avec sa cause probable et le correctif à envisager ;
- Résidus de vérification : chaque donnée signalée par l'agent vérificateur, tour par tour
  (diagnostic.verification), pour comprendre residus_apres_verification.
Les onglets Échecs et Résidus demandent des rapports produits avec REPORT_DETAIL=diagnostic (défaut).
Aucune valeur détectée n'y figure : seulement des types, des numéros de ligne et des indicateurs.

Les fichiers lus et produits contiennent des données réelles : ils restent sur le poste.
"""

from __future__ import annotations

import difflib
import importlib.util
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))
from pseudo.mails.parse import _body_alternatives, _leaves  # noqa: E402

_spec = importlib.util.spec_from_file_location("inspecter_eml", ROOT / "scripts" / "inspecter-eml.py")
inspecter = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(inspecter)

from email import policy  # noqa: E402
from email.parser import BytesParser  # noqa: E402

OCR_THRESHOLD = 0.75  # OCR_MIN_CONFIDENCE par défaut
MOTIFS = [
    "piece_jointe_non_prise_en_charge", "trop_de_pieces_jointes", "mail_illisible",
    "residus_apres_verification", "entites_non_localisees", "ocr_faible_confiance", "manuscrit_detecte",
    "image_illisible", "image_resolution_insuffisante", "image_floue", "image_trop_grande",
    "format_non_supporte", "fichier_trop_volumineux", "traitement_images_desactive", "traitement_mails_desactive",
    "pdf_illisible", "pdf_vide", "pdf_chiffre", "trop_de_pages",
]


def mail_parts(eml: Path):
    """Pièces d'un mail dans l'ordre et selon les règles du service (numéros identiques à pieces_refusees)."""
    msg = BytesParser(policy=policy.default).parse(eml.open("rb"))
    body = msg.get_body(preferencelist=("html", "plain"))
    alternatives = {id(p) for p in _body_alternatives(msg)}
    parts = []
    for part in _leaves(msg):
        if part is body or id(part) in alternatives:
            continue
        if part.get_content_type() == "message/rfc822":
            continue  # mail joint : compté à part par le service
        data = part.get_payload(decode=True)
        if data is None:
            continue
        parts.append((len(parts) + 1, part, data))
    return parts


def original_of(json_path: Path) -> Path | None:
    """Fichier d'origine en quarantaine à côté de son motif."""
    if json_path.name == "motif.json":
        emls = sorted(json_path.parent.glob("*.eml"))
        return emls[0] if emls else None
    candidate = json_path.with_name(json_path.name[: -len(".motif.json")])
    return candidate if candidate.exists() else None


def ocr_pages(report: dict) -> list[tuple[int, float | None, int | None]]:
    """(page, confiance, lignes) des pages lues par OCR ; None si le rapport ne le donne pas."""
    extraction = (report.get("diagnostic") or {}).get("extraction")
    if extraction:
        return [(e["page"], e.get("confiance_ocr"), e.get("lignes")) for e in extraction if e.get("confiance_ocr") is not None]
    return [(p, None, None) for p in report.get("pages_ocr_faible_confiance", [])]


# Cause probable d'un échec de localisation -> correctif à envisager
CAUSES = {
    "valeur coupée sur 2 lignes": "Prompt : un fragment par ligne avec le même group (règle 7) ; ou recherche sur 2 lignes jointes dans le code",
    "accents ou ponctuation modifiés": "Insister sur la copie exacte dans le prompt ; ou ignorer accents, apostrophes et ponctuation dans la recherche",
    "numéro de ligne inconnu": "Réduire LLM_CHUNK_CHARS (blocs plus courts)",
    "valeur légèrement modifiée": "Baisser FUZZY_MIN_SCORE (ex. 80)",
    "valeur transformée ou inventée": "Exemples dans le référentiel des champs (date convertie en chiffres, nom reconstitué…)",
}


def failure_cause(e: dict) -> str:
    """Cause probable d'un échec, d'après les indicateurs du diagnostic (voir docs/exploitation.md)."""
    if e.get("trouvee_sur_2_lignes"):
        return "valeur coupée sur 2 lignes"
    if e.get("trouvee_sans_accents_ponctuation"):
        return "accents ou ponctuation modifiés"
    if e.get("ligne_connue") is False:
        return "numéro de ligne inconnu"
    if max(e.get("meilleur_score_ligne") or 0, e.get("meilleur_score_document") or 0) >= 75:
        return "valeur légèrement modifiée"
    return "valeur transformée ou inventée"


def pipeline_reports(report: dict, is_mail: bool, original: Path | None):
    """Rapports du pipeline d'un motif : (partie, rapport, motifs de la partie). Un mail en a un pour
    son corps et un par pièce jointe ; un document, un seul."""
    if not is_mail:
        yield (original.name if original else "document"), report, report.get("motifs") or []
        return
    if report.get("corps"):
        yield "corps du mail", report["corps"], report["corps"].get("motifs") or []
    for a in report.get("pieces_jointes") or []:
        yield f"pièce : {a.get('fichier')}", a.get("rapport") or {}, a.get("motifs") or []


def safe_name(name: str) -> str:
    """Nom de fichier utilisable sous Windows."""
    return "".join("_" if c in '<>:"/\\|?*' or ord(c) < 32 else c for c in name).strip() or "piece"


def main(base: Path) -> None:
    """Lit tous les motifs sous base et écrit la synthèse Excel."""
    detail, refused, ocr, failures, residues = [], [], [], [], []
    without_diagnostic = 0
    for jp in sorted(base.rglob("*.json")):
        try:
            data = json.loads(jp.read_text(encoding="utf-8-sig"))
        except (OSError, ValueError):
            continue
        rel = jp.parent.relative_to(base)
        type_folder = rel.parts[0] if rel.parts else ""
        is_mail = jp.name == "motif.json"
        original = original_of(jp)
        report = data.get("rapport") or {}
        motifs = data.get("motifs") or []
        attachments = report.get("pieces_jointes") or []
        detail.append({
            "type": type_folder, "dossier": str(rel), "fichier": original, "nature": "mail" if is_mail else "document",
            "etape": data.get("etape"), "motifs": ", ".join(motifs),
            "pj": len(attachments) or None,
            "pj_quarantaine": sum(1 for a in attachments if a.get("statut") == "quarantine") or None,
            "json": jp,
        })

        parts = None
        if is_mail and original is not None and ("piece_jointe_non_prise_en_charge" in motifs or "ocr_faible_confiance" in motifs):
            parts = mail_parts(original)

        # 1. pièces refusées : structure MIME de chacune
        if "piece_jointe_non_prise_en_charge" in motifs and parts is not None:
            by_index = {n: (p, d) for n, p, d in parts}
            for r in report.get("pieces_refusees", []):
                part, raw = by_index.get(r.get("piece"), (None, b""))
                refused.append({
                    "type": type_folder, "dossier": str(rel), "mail": original, "piece": r.get("piece"),
                    "extension": r.get("extension"), "motif": r.get("motif"),
                    "nom": (part.get_filename() if part is not None else None) or "AUCUN",
                    "mime": part.get_content_type() if part is not None else None,
                    "disposition": (part.get_content_disposition() if part is not None else None) or "-",
                    "cid": "oui" if part is not None and part.get("content-id") else "non",
                    "taille": len(raw), "reel": inspecter.signature(raw) if raw else None,
                })

        # 2. OCR faible : confiance par page, pièces extraites pour être ouvertes
        if "ocr_faible_confiance" in motifs:
            sources = attachments if is_mail else [{"fichier": original.name if original else None, "rapport": report,
                                                     "motifs": report.get("motifs", motifs)}]
            for a in sources:
                if "ocr_faible_confiance" not in (a.get("motifs") or []):
                    continue
                target = original if not is_mail else None
                if is_mail and parts is not None:
                    # Le rapport donne le nom publié (dédoublonné, caractères neutralisés) : à défaut
                    # du nom exact, la pièce au nom le plus proche.
                    wanted = a.get("fichier") or ""
                    names = {Path((part.get_filename() or "").replace("\\", "/")).name: raw for _, part, raw in parts}
                    match = wanted if wanted in names else next(iter(difflib.get_close_matches(wanted, names, 1, 0.5)), None)
                    if match is not None:
                        out = jp.parent / "pieces_ocr_faible" / safe_name(wanted)
                        out.parent.mkdir(exist_ok=True)
                        out.write_bytes(names[match])
                        target = out
                    else:
                        target = original  # pièce non retrouvée : le mail, à ouvrir dans la messagerie
                for page, conf, lines in ocr_pages(a.get("rapport") or {}):
                    # Une page sans texte (vide, photo) a une confiance de 0 mais ne compte pas pour le motif.
                    if lines == 0:
                        low = "page sans texte"
                    else:
                        low = "oui" if conf is None or conf < OCR_THRESHOLD else "non"
                    ocr.append({
                        "type": type_folder, "dossier": str(rel), "nature": "mail" if is_mail else "document",
                        "fichier": a.get("fichier"), "page": page, "lignes": lines, "confiance": conf,
                        "faible": low, "ouvrir": target,
                    })

        # 3 et 4. Échecs de localisation et résidus de vérification (diagnostic des rapports)
        for partie, part_report, part_motifs in pipeline_reports(report, is_mail, original):
            diag = part_report.get("diagnostic")
            if not diag:
                if "entites_non_localisees" in part_motifs or "residus_apres_verification" in part_motifs:
                    without_diagnostic += 1
                continue
            where = {"type": type_folder, "dossier": str(rel), "fichier": original, "partie": partie,
                     "motifs_partie": ", ".join(part_motifs)}
            rounds = diag.get("verification") or []
            steps = [("détection", diag.get("localisation") or {})]
            steps += [(f"vérification tour {t.get('tour')}", t.get("localisation") or {}) for t in rounds]
            for step, loc in steps:
                for e in loc.get("echecs") or []:
                    cause = failure_cause(e)
                    failures.append({
                        **where, "etape": step, "champ": e.get("type"), "source": e.get("source"),
                        "ligne": e.get("ligne"), "ligne_connue": e.get("ligne_connue"), "longueur": e.get("longueur"),
                        "mots": e.get("mots"), "score_ligne": e.get("meilleur_score_ligne"),
                        "score_doc": e.get("meilleur_score_document"), "deux_lignes": e.get("trouvee_sur_2_lignes"),
                        "accents": e.get("trouvee_sans_accents_ponctuation"), "cause": cause, "correctif": CAUSES[cause],
                    })
            for t in rounds:
                last = t.get("tour") == len(rounds)
                common = {**where, "tour": t.get("tour"), "tours": len(rounds), "dernier": "oui" if last else "non",
                          "nouveau": "oui" if t.get("a_trouve_du_nouveau") else "non",
                          "caracteres": t.get("nouveaux_caracteres_masques"),
                          "non_convergee": "oui" if last and t.get("a_trouve_du_nouveau") else "non"}
                items = t.get("residus") or []
                if not items:
                    residues.append({**common, "champ": "(aucun)", "ligne": None})
                for r in items:
                    residues.append({**common, "champ": r.get("type"), "ligne": r.get("ligne")})

    write_workbook(base / "synthese-quarantaine.xlsx", detail, refused, ocr, failures, residues, without_diagnostic)
    print(f"{len(detail)} motif(s), {len(refused)} pièce(s) refusée(s), {len(ocr)} page(s) OCR, "
          f"{len(failures)} échec(s) de localisation, {len(residues)} ligne(s) de résidus "
          f"-> {base / 'synthese-quarantaine.xlsx'}")
    if without_diagnostic:
        print(f"{without_diagnostic} rapport(s) concerné(s) sans section diagnostic (produits avant REPORT_DETAIL)")


# --------------------------------------------------------------------------- Excel

FONT = Font(name="Arial", size=10)
HEADER = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="305496")
TITLE = Font(name="Arial", size=12, bold=True)
LINK = Font(name="Arial", size=10, color="0563C1", underline="single")


def table(ws, row: int, headers: list[str], rows: list[list], widths: list[int] | None = None) -> int:
    """Écrit un tableau (en-tête + lignes) à partir de row ; renvoie la ligne suivante."""
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row, c, h)
        cell.font, cell.fill = HEADER, HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r, values in enumerate(rows, row + 1):
        for c, v in enumerate(values, 1):
            cell = ws.cell(r, c)
            if isinstance(v, Path):
                cell.value, cell.hyperlink, cell.font = v.name, v.resolve().as_uri(), LINK
            else:
                cell.value, cell.font = v, FONT
    for c, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    return row + len(rows) + 2


def write_workbook(path: Path, detail: list[dict], refused: list[dict], ocr: list[dict],
                   failures: list[dict], residues: list[dict], without_diagnostic: int) -> None:
    """Classeur : Synthèse (formules sur les autres onglets), Détail, Pièces non prises en charge,
    OCR faible, Échecs de localisation, Résidus de vérification."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Synthèse"

    d = wb.create_sheet("Détail")
    table(d, 1, ["Dossier de type", "Dossier", "Fichier d'origine", "Nature", "Étape", "Motifs",
                 "Pièces jointes", "Pièces en quarantaine", "Motif (json)"],
          [[x["type"], x["dossier"], x["fichier"], x["nature"], x["etape"], x["motifs"], x["pj"], x["pj_quarantaine"], x["json"]]
           for x in detail], [16, 34, 60, 11, 18, 60, 10, 12, 30])
    d.freeze_panes = "A2"
    d.auto_filter.ref = f"A1:I{len(detail) + 1}"
    n = len(detail) + 1

    p = wb.create_sheet("Pièces non prises en charge")
    refused.sort(key=lambda x: (x["mime"] or "", x["dossier"], x["piece"] or 0))
    table(p, 1, ["Dossier de type", "Dossier du mail", "Mail", "N° pièce", "Extension (service)", "Motif",
                 "Nom dans le mail", "Type MIME", "Disposition", "Content-ID", "Taille (octets)", "Format réel (1ers octets)"],
          [[x["type"], x["dossier"], x["mail"], x["piece"], x["extension"], x["motif"], x["nom"], x["mime"],
            x["disposition"], x["cid"], x["taille"], x["reel"]] for x in refused],
          [16, 30, 50, 8, 12, 20, 40, 26, 12, 10, 14, 18])
    p.freeze_panes = "A2"
    p.auto_filter.ref = f"A1:L{len(refused) + 1}"
    np_ = len(refused) + 1

    o = wb.create_sheet("OCR faible")
    ocr.sort(key=lambda x: (x["faible"] == "page sans texte", x["confiance"] is None,
                            x["confiance"] if x["confiance"] is not None else 0))
    table(o, 1, ["Dossier de type", "Dossier", "Nature", "Fichier", "Page", "Lignes", "Confiance OCR", "Sous le seuil", "Fichier à ouvrir"],
          [[x["type"], x["dossier"], x["nature"], x["fichier"], x["page"], x["lignes"], x["confiance"], x["faible"], x["ouvrir"]]
           for x in ocr], [16, 34, 11, 50, 7, 8, 13, 16, 50])
    for r in range(2, len(ocr) + 2):
        o.cell(r, 7).number_format = "0.000"
    o.freeze_panes = "A2"
    o.auto_filter.ref = f"A1:I{len(ocr) + 1}"
    no = len(ocr) + 1

    f = wb.create_sheet("Échecs de localisation")
    failures.sort(key=lambda x: (x["cause"], x["champ"] or "", x["dossier"]))
    table(f, 1, ["Dossier de type", "Dossier", "Fichier", "Partie", "Étape", "Type de champ", "Agent", "Ligne",
                 "Ligne connue", "Longueur", "Mots", "Score ligne", "Score document", "Trouvée sur 2 lignes",
                 "Trouvée sans accents/ponctuation", "Cause probable", "Correctif à envisager", "Motifs de la partie"],
          [[x["type"], x["dossier"], x["fichier"], x["partie"], x["etape"], x["champ"], x["source"], x["ligne"],
            _yes(x["ligne_connue"]), x["longueur"], x["mots"], x["score_ligne"], x["score_doc"], _yes(x["deux_lignes"]),
            _yes(x["accents"]), x["cause"], x["correctif"], x["motifs_partie"]] for x in failures],
          [14, 30, 40, 30, 18, 16, 9, 9, 9, 9, 7, 9, 10, 10, 12, 30, 60, 30])
    f.freeze_panes = "A2"
    f.auto_filter.ref = f"A1:R{len(failures) + 1}"
    nf = max(len(failures) + 1, 2)

    v = wb.create_sheet("Résidus de vérification")
    residues.sort(key=lambda x: (x["dossier"], x["partie"], x["tour"] or 0))
    table(v, 1, ["Dossier de type", "Dossier", "Fichier", "Partie", "Tour", "Tours effectués", "Dernier tour",
                 "Type de champ", "Ligne", "Le tour a trouvé du nouveau", "Caractères masqués en plus",
                 "Non convergé (dernier tour encore du nouveau)", "Motifs de la partie"],
          [[x["type"], x["dossier"], x["fichier"], x["partie"], x["tour"], x["tours"], x["dernier"], x["champ"],
            x["ligne"], x["nouveau"], x["caracteres"], x["non_convergee"], x["motifs_partie"]] for x in residues],
          [14, 30, 40, 30, 6, 9, 9, 16, 9, 12, 12, 16, 30])
    v.freeze_panes = "A2"
    v.auto_filter.ref = f"A1:M{len(residues) + 1}"
    nv = max(len(residues) + 1, 2)

    # Synthèse : uniquement des formules sur les onglets de détail
    ws["A1"] = "Synthèse des mises en quarantaine"
    ws["A1"].font = TITLE
    ws["A2"] = f"Seuil OCR : {OCR_THRESHOLD} (OCR_MIN_CONFIDENCE par défaut). Un fichier peut avoir plusieurs motifs."
    ws["A2"].font = FONT
    present = [m for m in MOTIFS if any(m in x["motifs"].split(", ") for x in detail)]
    present += sorted({m for x in detail for m in x["motifs"].split(", ") if m and m not in present})
    rows = []
    for i, m in enumerate(present):
        r = 5 + i
        rows.append([m,
                     f'=COUNTIFS(Détail!$F$2:$F${n},"*{m}*",Détail!$D$2:$D${n},"mail")',
                     f'=COUNTIFS(Détail!$F$2:$F${n},"*{m}*",Détail!$D$2:$D${n},"document")',
                     f"=B{r}+C{r}", f"=D{r}/$D${5 + len(present)}"])
    total_row = 5 + len(present)
    rows.append(["Fichiers en quarantaine", f'=COUNTIFS(Détail!$D$2:$D${n},"mail")',
                 f'=COUNTIFS(Détail!$D$2:$D${n},"document")', f"=B{total_row}+C{total_row}", None])
    row = table(ws, 4, ["Motif", "Mails", "Documents", "Total", "% des fichiers"], rows, [34, 12, 12, 12, 14])
    for r in range(5, total_row):
        ws.cell(r, 5).number_format = "0%"
    for c in range(1, 5):
        ws.cell(total_row, c).font = Font(name="Arial", size=10, bold=True)

    folders = sorted({x["type"] for x in detail})
    ws.cell(row, 1, "Par dossier de type (nombre de fichiers ayant le motif)").font = TITLE
    rows = [[m] + [f'=COUNTIFS(Détail!$F$2:$F${n},"*{m}*",Détail!$A$2:$A${n},"{f}")' for f in folders] for m in present]
    row = table(ws, row + 1, ["Motif"] + folders, rows)

    mimes = sorted({(x["mime"] or "?", x["reel"] or "?") for x in refused})
    if mimes:
        ws.cell(row, 1, "Pièces non prises en charge, par type MIME et format réel").font = TITLE
        rows = [[mi, re_, f"=COUNTIFS('Pièces non prises en charge'!$H$2:$H${np_},\"{mi}\",'Pièces non prises en charge'!$L$2:$L${np_},\"{re_}\")",
                 f"=COUNTIFS('Pièces non prises en charge'!$H$2:$H${np_},\"{mi}\",'Pièces non prises en charge'!$L$2:$L${np_},\"{re_}\",'Pièces non prises en charge'!$G$2:$G${np_},\"AUCUN\")"]
                for mi, re_ in mimes]
        row = table(ws, row + 1, ["Type MIME", "Format réel", "Pièces", "dont sans nom"], rows)

    if ocr:
        ws.cell(row, 1, "Confiance OCR des pages").font = TITLE
        rng = f"'OCR faible'!$G$2:$G${no}"
        flag = f"'OCR faible'!$H$2:$H${no}"
        rows = [["Plus basse (pages avec texte)", f"=_xlfn.MINIFS({rng},{flag},\"<>page sans texte\")"],
                ["Moyenne (pages avec texte)", f"=AVERAGEIFS({rng},{flag},\"<>page sans texte\")"],
                ["Pages sous le seuil", f"=COUNTIFS({flag},\"oui\")"],
                ["Pages entre 0,70 et le seuil", f"=COUNTIFS({rng},\">=0.7\",{rng},\"<{OCR_THRESHOLD}\",{flag},\"oui\")"],
                ["Pages sans texte (ignorées par le motif)", f"=COUNTIFS({flag},\"page sans texte\")"]]
        row = table(ws, row + 1, ["Indicateur", "Valeur"], rows)

    # Échecs de localisation : par cause, et par cause et type de champ
    ech = "'Échecs de localisation'"
    ws.cell(row, 1, "Échecs de localisation, par cause probable").font = TITLE
    rows = [[c, f'=COUNTIFS({ech}!$P$2:$P${nf},"{c}")', CAUSES[c]] for c in CAUSES]
    rows.append(["Total", f"=COUNTA({ech}!$P$2:$P${nf})", None])
    row = table(ws, row + 1, ["Cause probable", "Échecs", "Correctif à envisager"], rows)
    champs = sorted({x["champ"] for x in failures if x["champ"]})
    if champs:
        ws.cell(row, 1, "Échecs de localisation, par cause et type de champ").font = TITLE
        rows = [[c] + [f'=COUNTIFS({ech}!$P$2:$P${nf},"{c}",{ech}!$F$2:$F${nf},"{t}")' for t in champs] for c in CAUSES]
        row = table(ws, row + 1, ["Cause probable"] + champs, rows)

    # Résidus de vérification : par tour et par type de champ
    res = "'Résidus de vérification'"
    ws.cell(row, 1, "Résidus de vérification").font = TITLE
    tours = sorted({x["tour"] for x in residues if x["tour"]})
    rows = [[f"Tour {t}",
             f'=COUNTIFS({res}!$E$2:$E${nv},{t},{res}!$H$2:$H${nv},"<>(aucun)")',
             f'=COUNTIFS({res}!$E$2:$E${nv},{t},{res}!$G$2:$G${nv},"oui",{res}!$L$2:$L${nv},"oui",{res}!$H$2:$H${nv},"<>(aucun)")']
            for t in tours]
    row = table(ws, row + 1, ["Tour", "Résidus signalés", "dont au dernier tour d'une partie non convergée"], rows)
    types_res = sorted({x["champ"] for x in residues if x["champ"] and x["champ"] != "(aucun)"})
    if types_res:
        rows = [[t, f'=COUNTIFS({res}!$H$2:$H${nv},"{t}")',
                 f'=COUNTIFS({res}!$H$2:$H${nv},"{t}",{res}!$L$2:$L${nv},"oui")'] for t in types_res]
        row = table(ws, row, ["Type de champ", "Résidus (tous tours)", "dont au dernier tour d'une partie non convergée"], rows)
    if without_diagnostic:
        ws.cell(row, 1, f"{without_diagnostic} rapport(s) en entites_non_localisees ou residus_apres_verification "
                        "sans section diagnostic (produits avant REPORT_DETAIL=diagnostic) : absents de ces tableaux.").font = FONT
    wb.save(path)


def _yes(value) -> str | None:
    """oui / non / vide pour un booléen du diagnostic."""
    return None if value is None else ("oui" if value else "non")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(Path(sys.argv[1]))
