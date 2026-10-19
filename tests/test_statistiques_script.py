"""scripts/statistiques.py : calculs sur des fiches réelles (produites par le traitement), écran et Excel."""

import importlib.util
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from fakes import FakeBlobService, MemoryStore, event, run_chain
from pseudo.router import route_event
from test_classification import make_runtime
from test_mails import make_mail, runtime as mail_runtime
from test_pipeline import make_text_pdf, services
from test_statistics import settings  # noqa: F401 - fixture

SCRIPT = Path(__file__).resolve().parent.parent / "scripts" / "statistiques.py"
spec = importlib.util.spec_from_file_location("statistiques", SCRIPT)
statistiques = importlib.util.module_from_spec(spec)
spec.loader.exec_module(statistiques)


def fiches(settings) -> list[dict]:
    """Fiches d'une petite journée : un PDF publié (traité deux fois), un mail avec une pièce, un ZIP refusé."""
    rt = make_runtime(settings, services())
    run_chain(rt, "lot1/facture.pdf")
    run_chain(rt, "lot1/facture.pdf")  # retraitement
    rt.blob_service.put("incoming", "lot1/archive.zip", b"PK")
    route_event(event("lot1/archive.zip"), rt)
    mail = mail_runtime(settings, make_mail())
    run_chain(mail, "lot1/mail.eml")
    return list(rt.mapping_store.statistics.values()) + list(mail.mapping_store.statistics.values())


def table(tables, name) -> dict:
    """Tableau sous forme {libellé ou 1re colonne: ligne}."""
    headers, rows = tables[name]
    return {r[0]: dict(zip(headers, r)) for r in rows}


def test_summary_counts_documents_pages_and_quarantine(settings):
    rows = fiches(settings)
    summary = {k: v["Valeur"] for k, v in table(statistiques.compute(rows, "jour", {"ocr": 0, "entree": 0, "sortie": 0, "language": 0}), "Synthèse").items()}
    assert summary["Documents déposés (un mail = 1)"] == 4  # 2 traitements du PDF + ZIP + mail
    assert summary["dont mails"] == 1 and summary["Documents en quarantaine"] == 1
    assert summary["% de mise en quarantaine"] == 25.0
    assert summary["Documents traités plus d'une fois"] == 1
    assert summary["Pages traitées (documents et pièces)"] == 3  # 2 x 1 page + 1 page de pièce
    assert summary["Pages passées par l'OCR (Document Intelligence)"] == 3
    assert summary["Pages passées par les agents IA"] == 3  # le ZIP refusé n'y passe pas
    assert summary["Textes passés par les agents IA (corps de mail, JSON)"] == 0  # agents simulés : pas d'appel compté
    assert summary["Coût estimé"] is None  # pas de prix


def test_by_type_and_rounds_and_cost(settings):
    rows = fiches(settings)
    prices = {"ocr": 1.5, "entree": 0.25, "sortie": 2.0, "language": 1.0}
    tables = statistiques.compute(rows, "mois", prices)
    by_type = {(r[0], r[1]): r for r in tables["Par type"][1]}
    assert by_type[("PDF hors mail", "application/pdf")][2] == 2
    assert by_type[("PDF dans un mail", "application/pdf")][2] == 1
    assert by_type[("Fichier refusé au routage", "application/zip")][4] == 1  # en quarantaine
    [rounds] = tables["Tours de vérification"][1]
    assert rounds[0] == "2026-03" and rounds[1] == 4  # 2 PDF, texte du mail, pièce
    summary = {k: v["Valeur"] for k, v in table(tables, "Synthèse").items()}
    assert summary["Coût estimé"] is not None and summary["Coût estimé"] >= 0


def test_main_prints_and_writes_the_workbook(settings, tmp_path, capsys):
    export = tmp_path / "fiches.json"
    export.write_text(json.dumps(fiches(settings)), encoding="utf-8")
    excel = tmp_path / "stats.xlsx"
    code = statistiques.main(["--fiches", str(export), "--debut", "2026-03-01", "--fin", "2026-03-31", "--excel", str(excel)])
    assert code == 0
    out = capsys.readouterr().out
    assert "% de mise en quarantaine" in out and "Tours de vérification" in out
    wb = load_workbook(excel)
    assert wb.sheetnames == ["Synthèse", "Par période", "Par type", "Tours de vérification", "Agents IA", "Documents"]
    assert wb["Documents"].max_row == 5  # en-tête + 4 documents


def test_cosmos_account_name_or_url():
    assert statistiques.cosmos_endpoint("cosmos-pseudo-dev-ab12") == "https://cosmos-pseudo-dev-ab12.documents.azure.com:443/"
    assert statistiques.cosmos_endpoint("https://x.documents.azure.com:443/") == "https://x.documents.azure.com:443/"


def _row(jour, depose, traite, categorie="pdf", niveau="document", statut="ok", job="j", **extra):
    """Fiche minimale pour les calculs de délai."""
    return {"id": job, "jobId": job, "documentId": job, "jour": jour, "hour": 10, "niveau": niveau,
            "categorie": categorie, "mime": "application/pdf", "statut": statut, "pages": 1,
            "deposeLe": depose, "traiteLe": traite, **extra}


def test_end_to_end_delay():
    rows = [
        _row("2026-10-18", "2026-10-18T08:00:00+00:00", "2026-10-18T08:00:30+00:00", job="a"),           # 30 s
        _row("2026-10-18", "2026-10-18T08:00:00+00:00", "2026-10-18T08:01:30+00:00", job="b",
             statut="quarantine"),                                                                         # 90 s
        _row("2026-10-18", "2026-10-18T08:00:00+00:00", "2026-10-18T08:10:00+00:00", job="m", categorie="mail"),  # 600 s
        # Pièce du mail, terminée avant lui : son délai ne compte pas à part.
        _row("2026-10-18", "2026-10-18T08:00:00+00:00", "2026-10-18T08:05:00+00:00", job="m-p1",
             categorie="pdf_mail", niveau="piece", mailJobId="m"),
        # Date de fin illisible : document compté, délai inconnu.
        _row("2026-10-19", "2026-10-19T08:00:00+00:00", "bad", job="z"),
    ]
    tables = statistiques.compute(rows, "jour", {"ocr": 0, "entree": 0, "sortie": 0, "language": 0})
    summary = {k: v["Valeur"] for k, v in table(tables, "Synthèse").items()}
    assert summary["Délai de bout en bout, moyenne (s)"] == 240.0  # (30 + 90 + 600) / 3 ; date illisible ignorée
    assert summary["Délai de bout en bout, médiane (s)"] == 90.0
    assert summary["Délai de bout en bout, maximum (s)"] == 600.0
    assert 90 < summary["Délai de bout en bout, 95e centile (s)"] <= 600
    period = table(tables, "Par période")["2026-10-18"]
    assert period["Délai moyen (s)"] == 240.0
    by_type = {(r[0], r[1]): dict(zip(tables["Par type"][0], r)) for r in tables["Par type"][1]}
    assert by_type[("PDF hors mail", "application/pdf")]["Délai moyen (s)"] == 60.0
    assert by_type[("PDF dans un mail", "application/pdf")]["Délai moyen (s)"] is None  # livré avec le mail
    docs = {r[2]: dict(zip(tables["Documents"][0], r)) for r in tables["Documents"][1]}
    assert docs["m"]["Délai (s)"] == 600.0 and docs["z"]["Délai (s)"] is None


def test_delay_of_a_single_document_and_of_none():
    assert statistiques.delay_summary([12.0])["p95"] == 12.0
    assert statistiques.delay_summary([None])["moyenne"] is None


def test_hours_narrow_the_period_and_group_by_hour(tmp_path, capsys):
    rows = [
        _row("2026-10-18", "2026-10-18T06:00:00+00:00", "2026-10-18T06:00:10+00:00", job=f"h{h}", hour=h)
        for h in (8, 9, 12, 13)
    ] + [_row("2026-10-19", "2026-10-19T06:00:00+00:00", "2026-10-19T06:00:10+00:00", job="next", hour=9)]
    assert {r["jobId"] for r in statistiques.in_hours(rows, "2026-10-18", 9, "2026-10-18", 12)} == {"h9", "h12"}
    assert {r["jobId"] for r in statistiques.in_hours(rows, "2026-10-18", 13, "2026-10-19", 9)} == {"h13", "next"}
    export = tmp_path / "fiches.json"
    export.write_text(json.dumps(rows), encoding="utf-8")
    excel = tmp_path / "stats.xlsx"
    statistiques.main(["--fiches", str(export), "--debut", "2026-10-18", "--fin", "2026-10-18",
                       "--heure-debut", "9", "--heure-fin", "12", "--par", "heure", "--excel", str(excel)])
    out = capsys.readouterr().out
    assert "2026-10-18 09h au 2026-10-18 12h (2 fiche(s)" in out
    periods = [row[0].value for row in load_workbook(excel)["Par période"].iter_rows(min_row=2)]
    assert periods == ["2026-10-18 09h", "2026-10-18 12h"]


def test_hours_are_checked():
    with pytest.raises(SystemExit):
        statistiques.parse_args(["--fiches", "x.json", "--heure-debut", "24"])
    with pytest.raises(SystemExit):
        statistiques.parse_args(["--fiches", "x.json", "--debut", "2026-10-18", "--fin", "2026-10-18",
                                 "--heure-debut", "14", "--heure-fin", "9"])
